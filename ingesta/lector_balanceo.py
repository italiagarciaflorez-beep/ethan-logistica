# ============================================================
#  ETHAN — Lector de archivos de balanceo v9
# ============================================================

import os
import re
import pandas as pd
from pathlib import Path
from datetime import datetime


def extraer_metadatos_nombre(ruta_archivo: str) -> dict:
    nombre = Path(ruta_archivo).stem.upper()

    marca = "desconocido"
    if "CRML" in nombre or "CARMEL" in nombre:
        marca = "carmel"
    elif "LGIN" in nombre or "LOGUIN" in nombre:
        marca = "loguin"
    elif "PCFK" in nombre or "PCKF" in nombre or "PACIFIKA" in nombre:
        marca = "pcfk"
    elif "DIGITAL" in nombre:
        marca = "digital"
    elif "ECUADOR" in nombre or "EXPO" in nombre:
        marca = "ecuador"

    campaña_str = "desconocida"
    numero_campaña = 0
    match = re.search(r'C[-_\s]?(\d+)', nombre)
    if match:
        numero_campaña = int(match.group(1))
        campaña_str = f"C-{numero_campaña:02d}"

    bin_str = "desconocido"
    match_bin = re.search(r'BIN[-_\s]?(\d+)', nombre)
    if match_bin:
        bin_str = f"BIN_{match_bin.group(1)}"

    linea_str = "desconocida"
    match_linea = re.search(r'LINEA[-_\s]?(\d+)', nombre)
    if match_linea:
        linea_str = f"LINEA_{int(match_linea.group(1)):02d}"

    fecha_mod = os.path.getmtime(ruta_archivo)
    fecha_mod_str = datetime.fromtimestamp(fecha_mod).strftime("%Y-%m-%d %H:%M")

    ruta_upper = str(ruta_archivo).upper()
    nivel = "historico" if "VIEJAS" in ruta_upper else "actual"

    return {
        "archivo":        Path(ruta_archivo).name,
        "ruta":           str(ruta_archivo),
        "marca":          marca,
        "campaña":        campaña_str,
        "numero_campaña": numero_campaña,
        "bin":            bin_str,
        "linea":          linea_str,
        "fecha_mod":      fecha_mod,
        "fecha_mod_str":  fecha_mod_str,
        "nivel":          nivel,
        "tipo":           "balanceo",
    }


def leer_hoja_analisis(ruta_archivo: str) -> dict:
    resultado = {
        "texto_completo":   "",
        "plu_total":        0,
        "estaciones_nec":   0,
        "estaciones_disp":  0,
        "ubicaciones":      0,
        "total_posiciones": 0,
        "tabla_resumen":    "",
    }

    try:
        df_analisis = pd.read_excel(ruta_archivo, sheet_name="Analisis", header=None)
        resultado["texto_completo"]   = df_analisis.to_string()
        resultado["plu_total"]        = _extraer_numero(df_analisis, ["Total plu", "Total de posiciones", "# Total"])
        resultado["estaciones_nec"]   = _extraer_numero(df_analisis, ["Estaciones Necesarias", "# Estaciones Necesarias"])
        resultado["estaciones_disp"]  = _extraer_numero(df_analisis, ["Estaciones disponibles", "# Estaciones disponibles"])
        resultado["ubicaciones"]      = _extraer_numero(df_analisis, ["# Ubicaciones", "Ubicaciones x estacion"])
        resultado["total_posiciones"] = _extraer_numero(df_analisis, ["# Total de posiciones", "Total de posiciones"])
    except Exception:
        pass

    tabla_generada = _generar_tabla_resumen(ruta_archivo)
    if tabla_generada:
        resultado["tabla_resumen"] = tabla_generada["tabla_str"]
        if resultado["plu_total"] == 0:
            resultado["plu_total"]        = tabla_generada["plu_total"]
            resultado["total_posiciones"] = tabla_generada["plu_total"]
        if resultado["texto_completo"] == "":
            resultado["texto_completo"] = tabla_generada["tabla_str"]

    return resultado


def _generar_tabla_resumen(ruta_archivo: str) -> dict | None:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(ruta_archivo, read_only=True, data_only=True)
        nombre_hoja = None
        for hoja in wb.sheetnames:
            if hoja.strip().upper() == "BALANCEO OFICIAL":
                nombre_hoja = hoja
                break
        wb.close()

        if not nombre_hoja:
            return None

        df = pd.read_excel(ruta_archivo, sheet_name=nombre_hoja)
        df.columns = [str(c).strip().upper() for c in df.columns]

        col_estacion  = _buscar_columna(df, ["ESTACION", "ESTACIÓN", "EST"])
        col_item      = _buscar_columna(df, ["ITEM/SAP", "ITEM_SAP", "ITEM SAP", "SAP"])
        col_tendencia = _buscar_columna(df, ["TENDENCIA", "TEND"])

        if not col_estacion:
            return None

        df = df.dropna(subset=[col_estacion])

        # CORRECCIÓN: Eliminar filas de TOTAL/GENERAL antes de agrupar
        df = df[~df[col_estacion].astype(str).str.upper().str.contains("TOTAL|GENERAL", na=False)]

        df[col_estacion] = pd.to_numeric(df[col_estacion], errors="coerce").fillna(0).astype(int)

        if col_tendencia:
            df[col_tendencia] = pd.to_numeric(df[col_tendencia], errors="coerce").fillna(0)

        agg = {col_item: "count"} if col_item else {}
        if col_tendencia:
            agg[col_tendencia] = "sum"

        if agg:
            tabla = df.groupby(col_estacion).agg(agg).reset_index()
        else:
            tabla = df.groupby(col_estacion).size().reset_index(name="count")

        nuevas_cols = ["Estación"]
        if col_item:
            nuevas_cols.append("Referencias (ITEM/SAP)")
        if col_tendencia:
            nuevas_cols.append("Tendencia")
        tabla.columns = nuevas_cols
        tabla = tabla.sort_values("Estación").reset_index(drop=True)

        total_ref  = int(tabla["Referencias (ITEM/SAP)"].sum()) if "Referencias (ITEM/SAP)" in tabla else 0
        total_tend = int(tabla["Tendencia"].sum()) if "Tendencia" in tabla else 0

        fila_total = {"Estación": "TOTAL GENERAL"}
        if "Referencias (ITEM/SAP)" in tabla.columns:
            fila_total["Referencias (ITEM/SAP)"] = total_ref
        if "Tendencia" in tabla.columns:
            fila_total["Tendencia"] = total_tend

        tabla = pd.concat([tabla, pd.DataFrame([fila_total])], ignore_index=True)

        for col in ["Referencias (ITEM/SAP)", "Tendencia"]:
            if col in tabla.columns:
                tabla[col] = tabla[col].apply(lambda x: f"{int(x):,}")

        return {
            "tabla_str":       tabla.to_string(index=False),
            "plu_total":       total_ref,
            "tendencia_total": total_tend,
        }

    except Exception as e:
        print(f"  No se pudo calcular tabla resumen: {e}")
        return None


def _buscar_columna(df: pd.DataFrame, opciones: list) -> str | None:
    for col in df.columns:
        for opcion in opciones:
            if opcion in col.upper():
                return col
    return None


def leer_hoja_oficial(ruta_archivo: str) -> pd.DataFrame:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(ruta_archivo, read_only=True, data_only=True)
        nombre_hoja = None
        for hoja in wb.sheetnames:
            if hoja.strip().upper() == "BALANCEO OFICIAL":
                nombre_hoja = hoja
                break
        wb.close()

        if not nombre_hoja:
            return pd.DataFrame()

        df = pd.read_excel(ruta_archivo, sheet_name=nombre_hoja)

        # Forzar nombre de columna D como descripcion
        if df.shape[1] >= 4:
            cols = list(df.columns)
            cols[3] = "DESCRIPCION"
            df.columns = cols

        df.columns = [str(c).strip().upper().replace(" ", "_") for c in df.columns]
        df = df.dropna(how="all")
        return df

    except Exception as e:
        print(f"Error leyendo BALANCEO OFICIAL: {e}")
        return pd.DataFrame()

def _extraer_numero(df: pd.DataFrame, etiquetas: list) -> int:
    texto = df.to_string().lower()
    for etiqueta in etiquetas:
        idx = texto.find(etiqueta.lower())
        if idx > -1:
            match = re.search(r'\d+', texto[idx:idx+100])
            if match:
                return int(match.group())
    return 0


def construir_texto_indexable(metadatos: dict, analisis: dict) -> str:
    from core.marcas import nombre_display
    marca_display = nombre_display(metadatos["marca"])

    texto = f"""
BALANCEO DE PICKING — {marca_display}
Archivo: {metadatos["archivo"]}
Marca: {marca_display}
Campaña: {metadatos["campaña"]}
Línea: {metadatos["linea"]}
BIN: {metadatos["bin"]}
Fecha de modificación: {metadatos["fecha_mod_str"]}
Tipo: {metadatos["nivel"].upper()} ({"más reciente" if metadatos["nivel"] == "actual" else "histórico"})

RESUMEN DEL BALANCEO:
Total PLU (referencias): {analisis.get("plu_total", "N/A")}
Estaciones necesarias: {analisis.get("estaciones_nec", "N/A")}
Estaciones disponibles: {analisis.get("estaciones_disp", "N/A")}
Ubicaciones por estación: {analisis.get("ubicaciones", "N/A")}
Total posiciones: {analisis.get("total_posiciones", "N/A")}

TABLA RESUMEN POR ESTACIÓN:
{analisis.get("tabla_resumen", "No disponible")}
""".strip()

    return texto